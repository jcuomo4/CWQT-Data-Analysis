#!/usr/bin/env python3
"""
CWQT (Community Water Quality Testing) maps & charts
=====================================================

Reads the CWQT master data workbook (downloaded from Google Sheets as .xlsx)
and produces maps and charts of Enterococcus results.

    python cwqt_visualize.py CWQT_Master_Data_Sheet.xlsx
    python cwqt_visualize.py CWQT_Master_Data_Sheet.xlsx --out figures --window-end 2026-08-31
    python cwqt_visualize.py CWQT_Master_Data_Sheet.xlsx --site-pngs      # also one PNG per site
    python cwqt_visualize.py --sheet-id 1AbC...xyz     # download straight from Google Sheets

Download the sheet from Google Sheets with File > Download > Microsoft Excel (.xlsx).
Thresholds, the wet-weather definition and the season dates are in the CONFIGURATION
section below.

Outputs (written to ./cwqt_outputs by default)
----------------------------------------------
  cwqt_map.html                      Interactive map. One layer per season plus an
                                     "all seasons" layer and a "latest 30 days" layer.
                                     Click a site for its stats and a mini trend chart.
  1_program_trend.png                Program-wide water quality by season, wet vs dry weather
  2_site_by_season_heatmap.png       Every site x every season, colored by seasonal geomean
  3_rainfall_effect.png              How results change with rainfall in the prior ~72 hours
  4_waterbody_comparison.png         Recent geomean by waterbody, one dot per site
  5_latest_30_day_status.png         Current 30-day geomean per site (what the GeoMean tabs show)
  site_timeseries.pdf                One page per site: every regular-season sample, by season
  tables/*.csv                       The numbers behind every chart (open in Sheets/Excel)
  data_quality_report.txt            Things in the workbook worth fixing

Requirements:  pip install pandas openpyxl matplotlib
               (plus google-auth requests only if downloading a private sheet)

How values are treated (matches the workbook FAQ and the GeoMean tabs)
----------------------------------------------------------------------
  * "<10"     -> 10      (non-detect, shown as the detection limit)
  * ">24196"  -> 24196   (above IDEXX upper limit with 10 ml dilution)
  * ">6000"   -> 6000    (same idea, other dilution)
  * "Trace" precipitation -> 0.0 in;  "0.45 (3.2 snow)" -> 0.45 in
  * Geometric mean = exp(mean(ln(MPN)))
"""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import sys
import warnings
from dataclasses import dataclass
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

import matplotlib
import numpy as np
import pandas as pd

matplotlib.use("Agg")  # draw to files, no window needed
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.backends.backend_pdf import PdfPages  # noqa: E402
from matplotlib.colors import ListedColormap  # noqa: E402
from matplotlib.lines import Line2D  # noqa: E402
from matplotlib.patches import Patch  # noqa: E402
from matplotlib.ticker import FixedLocator, FuncFormatter, NullLocator  # noqa: E402
from matplotlib.transforms import offset_copy  # noqa: E402

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# =============================================================================
# CONFIGURATION: edit these to change thresholds, colors or definitions
# =============================================================================

# Sheet / column names in the workbook. Columns are found by these names, so
# moving or inserting columns is fine; renaming them is not.
DATA_SHEET = "Data"
INFO_SHEET = "Information"
ALT_NAMES_SHEET = "Site Names + Site IDs"

COL_SITE = "Site"
COL_SITE_ID = "Site ID"
COL_DATE = "Full Date"
COL_TIME = "Sample Time"
COL_MPN = "Most Probable Number (MPN) of Enterococcus colonies per 100 ml"
COL_NOTES = "Notes"
# Daily rainfall columns (inches, Central Park), most recent day first.
# "Wet weather" = sum of the first RAIN_DAYS of these >= WET_WEATHER_INCHES.
RAIN_COLS = [
    "Day of Collection Precipitation (Thursday)",
    "Previous Day Precipitation (Wednesday)",
    "Previous Tuesday Precipitation",
    "Previous Monday Precipitation",
    "Previous Sunday Precipitation",
    "Previous Saturday Precipitation",
    "Previous Friday Precipitation",
]
RAIN_DAYS = 3                # day of sampling + 2 previous days (~72 hours)
WET_WEATHER_INCHES = 0.25

# Water quality categories (NYC DOH primary contact, the same green/yellow/red
# scheme used on the "30 Day GeoMean Primary Contact" tab).
GOOD_MAX = 35                # geomean <= 35  -> acceptable for primary contact
ADVISORY_MAX = 104           # 35-104 -> advisory;  > 104 -> closure-level
SINGLE_SAMPLE_LIMIT = 104    # single-sample threshold used for "% of samples over"
MIN_SAMPLES_FOR_GEOMEAN = 5  # standards call for >= 5 samples in a 30-day geomean

# NYS DEC class-specific criteria (geomean, statistical threshold value), shown
# in map popups. Class I numbers are DEC's 72/266 criteria used on the Class I tab.
DEC_CRITERIA = {"SB": (35, 130), "SA": (35, 130), "I": (72, 266)}

# How many recent seasons to use for "recent average" charts (waterbody chart).
RECENT_SEASONS = 5

# Regular sampling season (month, day), inclusive. Seasonal statistics use only
# samples in this window; winter / post-season samples are kept out of them.
# (The first Thursday of October has usually been the last regular-season test.)
SEASON_START = (5, 1)
SEASON_END = (10, 5)

# "Long-running sites" on the trend chart: sites with a seasonal geomean
# (>= MIN_SAMPLES_FOR_GEOMEAN samples) in at least this many seasons.
CORE_MIN_SEASONS = 10

# Colors. Status colors are reserved for good / advisory / poor and are always
# paired with a text label. SERIES is for ordinary data marks.
COLOR_GOOD = "#0ca30c"
COLOR_ADVISORY = "#fab219"
COLOR_POOR = "#d03b3b"
COLOR_NODATA = "#c3c2b7"
SERIES = "#2a78d6"          # blue  - main series
SERIES_2 = "#eb6834"        # orange - second series (wet weather)
INK = "#0b0b0b"
INK_2 = "#52514e"
MUTED = "#898781"
GRID = "#e1e0d9"
SURFACE = "#fcfcfb"

CATEGORY_LABELS = {
    "good": f"Good: geomean ≤ {GOOD_MAX}",
    "advisory": f"Advisory: {GOOD_MAX}–{ADVISORY_MAX}",
    "poor": f"Poor: > {ADVISORY_MAX}",
    "nodata": "No data",
}
CATEGORY_COLORS = {"good": COLOR_GOOD, "advisory": COLOR_ADVISORY,
                   "poor": COLOR_POOR, "nodata": COLOR_NODATA}


# =============================================================================
# LOADING & CLEANING
# =============================================================================

@dataclass
class QAIssue:
    category: str
    detail: str


def _norm(s) -> str:
    """Normalize a site name for matching (case, spaces, punctuation)."""
    s = str(s).lower().strip()
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def parse_mpn(v):
    """Return (numeric_value, censor_flag) for one MPN cell.

    censor_flag is '<' for non-detects, '>' for above-range, '' otherwise.
    """
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return np.nan, ""
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v), ""
    s = str(v).strip().replace(",", "")
    m = re.fullmatch(r"([<>])?\s*(\d+(?:\.\d+)?)", s)
    if not m:
        return np.nan, "?"
    return float(m.group(2)), (m.group(1) or "")


def parse_rain(v) -> float:
    """'Trace' -> 0, '0.45 (3.2 snow)' -> 0.45, numbers pass through."""
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return np.nan
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v)
    s = str(v).strip().lower()
    if s.startswith("trace") or s in {"t", "tr"}:
        return 0.0
    m = re.match(r"(\d*\.?\d+)", s)
    return float(m.group(1)) if m else np.nan


def _find_header_row(raw: pd.DataFrame, must_contain: str) -> int:
    for i in range(min(len(raw), 40)):
        if any(str(x).strip() == must_contain for x in raw.iloc[i].tolist()):
            return i
    raise ValueError(f"Could not find a header row containing '{must_contain}'")


def load_sites(xlsx: Path, qa: list[QAIssue]) -> pd.DataFrame:
    """Site table from the Information tab (one row per site, with coordinates)."""
    raw = pd.read_excel(xlsx, sheet_name=INFO_SHEET, header=None)
    hdr = _find_header_row(raw, "Latitude")
    cols = [str(c).strip() for c in raw.iloc[hdr].tolist()]
    body = raw.iloc[hdr + 1:].copy()
    body.columns = cols

    # The tab has two "Site ID" columns; use the one immediately after the name.
    name_idx = next(i for i, c in enumerate(cols) if c.startswith("Table of Contents - Site"))
    sites = pd.DataFrame({
        "site": body.iloc[:, name_idx],
        "site_id": body.iloc[:, name_idx + 1],
    })
    keep = {
        "Currently Testing": "currently_testing",
        "Waterbody": "waterbody",
        "Borough/State": "borough",
        "Neighborhood/Town": "neighborhood",
        "Latitude": "lat",
        "Longitude": "lon",
        "NYS DEC Water Body Classification": "dec_class",
    }
    for src, dst in keep.items():
        idx = cols.index(src)
        sites[dst] = body.iloc[:, idx].values
    sites = sites.dropna(subset=["site"])
    sites = sites[pd.to_numeric(sites["site_id"], errors="coerce").notna()].copy()
    sites["site_id"] = sites["site_id"].astype(int)
    sites["site"] = sites["site"].astype(str).str.strip()
    sites["lat"] = pd.to_numeric(sites["lat"], errors="coerce")
    sites["lon"] = pd.to_numeric(sites["lon"], errors="coerce")
    sites["currently_testing"] = sites["currently_testing"].astype(str).str.strip().str.lower().eq("x")
    sites["dec_class"] = sites["dec_class"].fillna("").astype(str).str.strip()
    sites["waterbody"] = sites["waterbody"].fillna("Unknown").astype(str).str.strip()

    # Coordinate sanity checks (NYC area: lat ~40.4-41.2, lon ~ -74.6 to -73.6)
    for _, r in sites.iterrows():
        if pd.notna(r.lon) and r.lon > 0:
            qa.append(QAIssue("Coordinates",
                f"Site {r.site_id} '{r.site}': longitude {r.lon} is positive "
                f"(missing minus sign). Using {-r.lon} for maps."))
        if pd.isna(r.lat) or pd.isna(r.lon):
            qa.append(QAIssue("Coordinates", f"Site {r.site_id} '{r.site}' has no coordinates."))
    sites.loc[sites["lon"] > 0, "lon"] *= -1

    dups = sites[sites.site_id.duplicated(keep=False)]
    for sid, g in dups.groupby("site_id"):
        qa.append(QAIssue("Site IDs", f"Site ID {sid} is used for: {'; '.join(g.site)}"))
    return sites.drop_duplicates("site_id").reset_index(drop=True)


def _name_to_id_lookup(xlsx: Path, sites: pd.DataFrame) -> dict[str, int]:
    lookup = {_norm(n): i for n, i in zip(sites.site, sites.site_id)}
    try:
        alt = pd.read_excel(xlsx, sheet_name=ALT_NAMES_SHEET)
        id_col = alt.columns[0]
        for _, r in alt.iterrows():
            if pd.isna(r[id_col]):
                continue
            for c in alt.columns[1:]:
                if pd.notna(r[c]):
                    lookup.setdefault(_norm(r[c]), int(r[id_col]))
    except (ValueError, KeyError):
        pass
    return lookup


def load_samples(xlsx: Path, sites: pd.DataFrame, qa: list[QAIssue],
                 average_splits: bool = False) -> pd.DataFrame:
    """One row per sample, cleaned. Values come from the workbook's last
    calculated values (Google Sheets includes these when you download .xlsx)."""
    d = pd.read_excel(xlsx, sheet_name=DATA_SHEET)
    d.columns = [str(c).strip() for c in d.columns]
    missing = [c for c in [COL_SITE, COL_DATE, COL_MPN] if c not in d.columns]
    if missing:
        sys.exit(f"Data tab is missing expected column(s): {missing}")
    d = d.dropna(subset=[COL_SITE, COL_DATE], how="any").copy()

    out = pd.DataFrame({
        "site_name_in_data": d[COL_SITE].astype(str).str.strip(),
        "date": pd.to_datetime(d[COL_DATE], errors="coerce"),
        "raw_mpn": d[COL_MPN],
        "notes": d.get(COL_NOTES, pd.Series(index=d.index, dtype=object)),
        "row": d.index + 2,  # spreadsheet row number, for QA messages
    })

    # Site ID: use the sheet's value; if blank/error, look the name up.
    lookup = _name_to_id_lookup(xlsx, sites)
    sid = pd.to_numeric(d.get(COL_SITE_ID), errors="coerce")
    by_name = out["site_name_in_data"].map(lambda n: lookup.get(_norm(n)))
    out["site_id"] = sid.fillna(by_name)
    unresolved = out["site_id"].isna()
    for n, g in out[unresolved].groupby("site_name_in_data"):
        qa.append(QAIssue("Site IDs", f"'{n}' ({len(g)} rows, e.g. row {g.row.iloc[0]}) "
                                      "has no Site ID and no matching name in Information."))
    out = out[~unresolved].copy()
    out["site_id"] = out["site_id"].astype(int)

    unknown = set(out.site_id) - set(sites.site_id)
    for s in sorted(unknown):
        qa.append(QAIssue("Site IDs", f"Site ID {s} appears in Data but not in Information."))

    # Name in Data vs. name in Information for the same ID
    info_name = sites.set_index("site_id")["site"]
    mism = out[out.site_id.map(info_name).map(_norm) != out.site_name_in_data.map(_norm)]
    for (sid_, n), g in mism.groupby(["site_id", "site_name_in_data"]):
        qa.append(QAIssue("Site names",
            f"Data row(s) {g.row.min()}–{g.row.max()} ({len(g)}): '{n}' vs Information "
            f"'{info_name.get(sid_, '?')}' (ID {sid_})."))

    # MPN
    parsed = out["raw_mpn"].map(parse_mpn)
    out["mpn"] = parsed.map(lambda t: t[0])
    out["censor"] = parsed.map(lambda t: t[1])
    for _, r in out[out.censor == "?"].iterrows():
        qa.append(QAIssue("MPN values", f"Row {r.row}: unreadable MPN '{r.raw_mpn}'."))
    blank = out[out.mpn.isna() & (out.censor != "?")]
    if len(blank):
        qa.append(QAIssue("MPN values",
            f"{len(blank)} row(s) have a date and site but no MPN result "
            f"(rows {', '.join(map(str, blank.row.head(10)))}"
            f"{'…' if len(blank) > 10 else ''})."))

    # Lab/field splits recorded in Notes, e.g. "<10 (lab split)"
    split_re = re.compile(r"^\s*([<>]?\s*\d+(?:\.\d+)?)\s*\((?:(?:lab|field)\s*split|duplicate)\)", re.I)
    split_val = out["notes"].map(lambda s: parse_mpn(split_re.match(str(s)).group(1))[0]
                                 if isinstance(s, str) and split_re.match(s) else np.nan)
    out["split_mpn"] = split_val
    if average_splits:  # FAQ: "Average the values"
        has = out.split_mpn.notna() & out.mpn.notna()
        out.loc[has, "mpn"] = (out.loc[has, "mpn"] + out.loc[has, "split_mpn"]) / 2

    # Rain
    rain_cols = [c for c in RAIN_COLS if c in d.columns]
    rain = pd.DataFrame({c: d[c].map(parse_rain) for c in rain_cols})
    use = rain_cols[:RAIN_DAYS]
    out["rain_in"] = rain[use].sum(axis=1, min_count=len(use))
    out["wet"] = np.where(out.rain_in.isna(), np.nan, out.rain_in >= WET_WEATHER_INCHES)

    out["year"] = out.date.dt.year
    md = out.date.dt.month * 100 + out.date.dt.day
    out["in_season"] = md.between(SEASON_START[0] * 100 + SEASON_START[1],
                                  SEASON_END[0] * 100 + SEASON_END[1])
    out = out.merge(sites[["site_id", "site", "waterbody", "borough", "dec_class",
                           "lat", "lon", "currently_testing"]], on="site_id", how="left")
    out = out.dropna(subset=["date", "mpn"]).sort_values(["site_id", "date"]).reset_index(drop=True)

    # Other QA observations
    wd = out.date.dt.dayofweek
    nthu = int((wd != 3).sum())
    if nthu:
        qa.append(QAIssue("Rainfall",
            f"{nthu} samples were not collected on a Thursday, but rainfall columns are "
            "labeled by weekday (e.g. 'Previous Day Precipitation (Wednesday)')."))
    no_rain = out[out.rain_in.isna()].groupby("year").size()
    for y, n in no_rain.items():
        qa.append(QAIssue("Rainfall", f"{y}: {n} samples have no rainfall data."))
    dupe = out[out.duplicated(["site_id", "date"], keep=False)]
    if len(dupe):
        groups = dupe.groupby(["site_id", "date"]).row.apply(lambda r: "/".join(map(str, r)))
        more = "…" if len(groups) > 10 else ""
        qa.append(QAIssue("Duplicates",
            f"{len(groups)} site+date combinations have more than one row (Data rows "
            f"{', '.join(groups.head(10))}{more}). Geomeans count each row."))
    qc_re = re.compile(r"incubat|CFU|sterile|processed \d|hours", re.I)
    qc = out[out.notes.astype(str).str.contains(qc_re) & out.notes.notna()]
    for note, g in qc.groupby("notes"):
        qa.append(QAIssue("Lab QC notes (results kept; consider a QC-flag column)",
            f"'{note}': {len(g)} row(s), e.g. row {g.row.min()} ({g.date.min():%Y-%m-%d})."))
    return out


# =============================================================================
# SUMMARY STATISTICS
# =============================================================================

def geomean(x) -> float:
    x = np.asarray(x, dtype=float)
    x = x[np.isfinite(x) & (x > 0)]
    return float(np.exp(np.log(x).mean())) if len(x) else np.nan


def category(gm: float) -> str:
    if gm is None or not np.isfinite(gm):
        return "nodata"
    if gm <= GOOD_MAX:
        return "good"
    if gm <= ADVISORY_MAX:
        return "advisory"
    return "poor"


def summarize(g: pd.DataFrame) -> pd.Series:
    wet = g[g.wet == 1]
    dry = g[g.wet == 0]
    return pd.Series({
        "n": len(g),
        "geomean": geomean(g.mpn),
        "pct_over_limit": 100 * (g.mpn > SINGLE_SAMPLE_LIMIT).mean(),
        "geomean_dry": geomean(dry.mpn),
        "geomean_wet": geomean(wet.mpn),
        "n_dry": len(dry),
        "n_wet": len(wet),
        "first": g.date.min(),
        "last": g.date.max(),
        "last_mpn": g.sort_values("date").mpn.iloc[-1],
    })


def site_stats(samples: pd.DataFrame, by_year: bool = False) -> pd.DataFrame:
    keys = ["site_id", "year"] if by_year else ["site_id"]
    s = samples.groupby(keys).apply(summarize, include_groups=False).reset_index()
    s["n"] = s["n"].astype(int)
    s["category"] = s.geomean.map(category)
    return s


def rolling_window_stats(samples: pd.DataFrame, end: pd.Timestamp, days: int = 30) -> pd.DataFrame:
    """Stats for each site over the `days`-day window ending on `end` (inclusive),
    the same calculation as the GeoMean tabs."""
    start = end - pd.Timedelta(days=days - 1)
    w = samples[(samples.date >= start) & (samples.date <= end)]
    s = site_stats(w)
    s["window_start"], s["window_end"] = start, end
    return s


# =============================================================================
# VERIFY AGAINST THE WORKBOOK'S OWN GEOMEAN TABS
# =============================================================================

def verify_against_workbook(xlsx: Path, samples: pd.DataFrame) -> list[str]:
    """Recompute a few numbers the workbook already shows and report agreement.
    Uses openpyxl cached values; silently skipped if the tab layout changes."""
    msgs = []
    try:
        import openpyxl
        wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
        ws = wb["30 Day GeoMean Primary Contact "] if "30 Day GeoMean Primary Contact " in wb.sheetnames \
            else wb["30 Day GeoMean Primary Contact"]
        rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=6, values_only=True))
        year = rows[0][0]
        start, end = rows[6][4], rows[6][5]  # E7, F7
        w = samples[(samples.date >= pd.Timestamp(start)) & (samples.date <= pd.Timestamp(end))
                    & (samples.year == year)]
        ytd = samples[samples.year == year]
        n_ok = n_tot = 0
        for r in rows[8:]:
            sid, cnt, ytd_gm, gm30 = r[1], r[2], r[3], r[4]
            if not isinstance(sid, (int, float)):
                continue
            ws_ = w[w.site_id == sid]
            mine = (len(ws_), geomean(ytd[ytd.site_id == sid].mpn), geomean(ws_.mpn))
            theirs = (cnt, ytd_gm, gm30)
            ok = True
            for a, b in zip(mine, theirs):
                if isinstance(b, (int, float)):
                    ok = ok and np.isfinite(a) and abs(a - b) < 0.01 * max(1, abs(b))
                else:  # the tab shows "No Data"
                    ok = ok and (a == 0 or not np.isfinite(a))
            n_ok += bool(ok)
            n_tot += 1
        msgs.append(f"30 Day GeoMean Primary Contact tab: {n_ok}/{n_tot} sites match "
                    f"(count, YTD geomean, 30-day geomean).")
    except Exception as e:  # pragma: no cover
        msgs.append(f"(Skipped workbook cross-check: {e})")
    return msgs


# =============================================================================
# CHARTS (matplotlib)
# =============================================================================


MIN_SAMPLES_PER_SEASON = 100   # seasons with fewer samples are left off the trend chart
MPN_FLOOR, MPN_CEIL = 7, 40000  # y-limits for MPN axes (log scale)
MPN_TICKS = [10, 100, 1000, 10000]


def _style():
    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": ["Helvetica Neue", "Helvetica", "Arial", "DejaVu Sans"],
        "font.size": 9.5,
        "text.color": INK,
        "axes.labelcolor": INK_2,
        "axes.labelsize": 9.5,
        "xtick.color": MUTED, "ytick.color": MUTED,
        "xtick.labelcolor": INK_2, "ytick.labelcolor": INK_2,
        "axes.edgecolor": "#c3c2b7", "axes.linewidth": 0.8,
        "axes.facecolor": SURFACE, "figure.facecolor": SURFACE, "savefig.facecolor": SURFACE,
        "axes.spines.top": False, "axes.spines.right": False,
        "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.7, "grid.linestyle": "-",
        "axes.axisbelow": True,
        "axes.titlesize": 12.5, "axes.titleweight": "bold", "axes.titlelocation": "left",
        "axes.titlepad": 10,
        "legend.frameon": False, "legend.fontsize": 9,
        "savefig.dpi": 160, "savefig.bbox": "tight", "savefig.pad_inches": 0.25,
        "pdf.fonttype": 42,
    })


def _fmt_int(v, _pos=None):
    return f"{int(round(v)):,}"


def _mpn_axis(ax, axis="y"):
    setter = ax.set_yscale if axis == "y" else ax.set_xscale
    setter("log")
    a = ax.yaxis if axis == "y" else ax.xaxis
    a.set_major_locator(FixedLocator(MPN_TICKS))
    a.set_minor_locator(NullLocator())
    a.set_major_formatter(FuncFormatter(_fmt_int))
    (ax.set_ylim if axis == "y" else ax.set_xlim)(MPN_FLOOR, MPN_CEIL)


def _quality_bands(ax, axis="y", tint=True, labels=True, label_pos=1.002):
    """Good / advisory / poor zones: light tints, dashed threshold lines, and
    zone labels placed outside the plot so they never collide with data labels.
    Call after the axis limits are set."""
    span = ax.axhspan if axis == "y" else ax.axvspan
    line = ax.axhline if axis == "y" else ax.axvline
    lo_lim, hi_lim = sorted(ax.get_ylim() if axis == "y" else ax.get_xlim())
    zones = [(lo_lim, GOOD_MAX, COLOR_GOOD, f"Good  ≤{GOOD_MAX}"),
             (GOOD_MAX, ADVISORY_MAX, COLOR_ADVISORY, f"Advisory  {GOOD_MAX}–{ADVISORY_MAX}"),
             (ADVISORY_MAX, hi_lim, COLOR_POOR, f"Poor  >{ADVISORY_MAX}")]
    for lo, hi, c, _ in zones:
        if tint:
            span(lo, hi, color=c, alpha=0.08, lw=0, zorder=0)
    for v in (GOOD_MAX, ADVISORY_MAX):
        line(v, color=INK_2, lw=0.9, ls=(0, (4, 3)), alpha=0.7, zorder=1)
    if not labels:
        return
    for lo, hi, c, txt in zones:
        mid = math.sqrt(lo * hi)
        if axis == "y":
            tr = ax.get_yaxis_transform()
            # a colored rail just outside the right edge, with the label beside it
            ax.plot([1.012, 1.012], [lo * 1.04, hi / 1.04], transform=tr, color=c, lw=4,
                    solid_capstyle="butt", clip_on=False, zorder=5)
            ax.text(1.026, mid, txt.replace("  ", "\n", 1), transform=tr, ha="left",
                    va="center", fontsize=8, color=INK_2, linespacing=1.1, clip_on=False)
        else:
            tr = ax.get_xaxis_transform()
            ax.text(mid, label_pos, txt, transform=tr, ha="center", va="bottom", fontsize=8,
                    color=INK_2)


def _subtitle(ax, text, y=1.0):
    ax.annotate(text, xy=(0, y), xycoords="axes fraction", xytext=(0, 6),
                textcoords="offset points", fontsize=9, color=INK_2, va="bottom", ha="left")


def short_name(site: str, waterbody: str) -> str:
    pre = f"{waterbody}, "
    return site[len(pre):] if site.startswith(pre) else site


# ---- 1. Program trend -------------------------------------------------------

def chart_program_trend(season, site_year, out: Path, tables: Path):
    per_season = season.groupby("year").size()
    years = [int(y) for y in per_season.index if per_season[y] >= MIN_SAMPLES_PER_SEASON]
    sy = site_year[(site_year.n >= MIN_SAMPLES_FOR_GEOMEAN) & site_year.year.isin(years)]
    seasons_per_site = sy.groupby("site_id").year.nunique()
    core = seasons_per_site[seasons_per_site >= CORE_MIN_SEASONS].index
    syc = sy[sy.site_id.isin(core)]

    agg = sy.groupby("year").geomean.agg(
        median_site_geomean="median",
        q25=lambda x: x.quantile(0.25), q75=lambda x: x.quantile(0.75), n_sites="size")
    agg["median_long_running_sites"] = syc.groupby("year").geomean.median()
    agg["n_long_running_sites"] = syc.groupby("year").size()
    s = season[season.year.isin(years)].assign(over=lambda x: x.mpn > SINGLE_SAMPLE_LIMIT)
    agg["pct_over_all"] = s.groupby("year").over.mean() * 100
    agg["pct_over_dry"] = s[s.wet == 0].groupby("year").over.mean() * 100
    agg["pct_over_wet"] = s[s.wet == 1].groupby("year").over.mean() * 100
    agg["n_samples"] = s.groupby("year").size()
    _r(agg).to_csv(tables / "program_trend_by_season.csv")

    fig, (a1, a2) = plt.subplots(2, 1, figsize=(11, 8.6), sharex=True,
                                 gridspec_kw={"height_ratios": [1.15, 1], "hspace": 0.36})
    x = agg.index.values
    _mpn_axis(a1)
    a1.set_ylim(MPN_FLOOR, 2000)
    _quality_bands(a1, tint=False)
    a1.fill_between(x, agg.q25, agg.q75, color=SERIES, alpha=0.10, lw=0, zorder=2)
    a1.plot(x, agg.median_site_geomean, color=SERIES, lw=2, marker="o", ms=5, mec=SURFACE,
            mew=1.5, zorder=4, label="All sites sampled that season (median; band = middle 50%)")
    a1.plot(x, agg.median_long_running_sites, color=MUTED, lw=1.8, marker="o", ms=4, mec=SURFACE,
            mew=1.2, zorder=3, label=f"Only the {len(core)} long-running sites (median)")
    a1.annotate(f"{agg.median_site_geomean.iloc[-1]:,.0f}",
                (x[-1], agg.median_site_geomean.iloc[-1]), xytext=(8, 0),
                textcoords="offset points", va="center", fontsize=9, color=INK, fontweight="bold")
    a1.set_ylabel("Seasonal geometric mean\n(MPN / 100 ml, log scale)")
    a1.set_title("Typical site water quality, by season", pad=28)
    _subtitle(a1, f"Each site's seasonal geomean needs ≥{MIN_SAMPLES_FOR_GEOMEAN} samples. "
                  f"Long-running = a seasonal geomean in ≥{CORE_MIN_SEASONS} seasons, so the "
                  f"gray line isn't affected by sites being added.")
    a1.legend(loc="upper left", fontsize=8.5, handlelength=1.6)
    a1.set_xlim(x.min() - 0.5, x.max() + 0.7)
    a1.grid(axis="x", visible=False)

    series = [("pct_over_all", "All samples", MUTED, 1.8),
              ("pct_over_dry", f"Dry weather (<{WET_WEATHER_INCHES}\" rain)", SERIES, 2),
              ("pct_over_wet", f"Wet weather (≥{WET_WEATHER_INCHES}\" rain)", SERIES_2, 2)]
    for col, lab, c, lw in series:
        v = agg[col].dropna()
        a2.plot(v.index, v.values, color=c, lw=lw, marker="o", ms=4.5, mec=SURFACE, mew=1.2,
                label=lab, zorder=3)
        if len(v):
            a2.annotate(f"{v.iloc[-1]:.0f}%", (v.index[-1], v.iloc[-1]), xytext=(7, 0),
                        textcoords="offset points", va="center", fontsize=8.5, color=INK_2)
    a2.set_ylim(0, 100)
    a2.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:.0f}%"))
    a2.set_ylabel(f"Samples over {SINGLE_SAMPLE_LIMIT} MPN")
    a2.set_title(f"Share of samples over the single-sample limit ({SINGLE_SAMPLE_LIMIT})", pad=28)
    rain_txt = (f"Wet weather = ≥{WET_WEATHER_INCHES}\" of rain at Central Park on the sampling "
                f"day and the {RAIN_DAYS - 1} days before.")
    no_rain = [y for y in years if pd.isna(agg.loc[y, "pct_over_dry"]) and pd.isna(agg.loc[y, "pct_over_wet"])]
    if no_rain:
        rain_txt += f" No rainfall entered yet for {', '.join(map(str, no_rain))}."
    _subtitle(a2, rain_txt)
    a2.legend(loc="upper left", ncols=3, fontsize=8.5, handlelength=1.6)
    a2.set_xticks(x)
    a2.set_xticklabels([f"{y}\n{int(agg.loc[y, 'n_sites'])}" for y in x], fontsize=8.5)
    a2.set_xlabel("Season  (number underneath = sites with a seasonal geomean)", fontsize=8.5)
    a2.grid(axis="x", visible=False)
    fig.savefig(out)
    plt.close(fig)


# ---- 2. Site x season heatmap ----------------------------------------------

def chart_site_heatmap(site_year, sites, out: Path, tables: Path):
    sy = site_year.merge(sites[["site_id", "site", "waterbody"]], on="site_id")
    years = list(range(int(sy.year.min()), int(sy.year.max()) + 1))
    wide = sy.pivot(index="site_id", columns="year", values="geomean").reindex(columns=years)
    nwide = sy.pivot(index="site_id", columns="year", values="n").reindex(columns=years)
    meta = sites.set_index("site_id").loc[wide.index, ["site", "waterbody"]]
    _r(meta.join(wide).sort_values(["waterbody", "site"])).to_csv(tables / "site_by_season_geomean.csv")

    rows, labels, is_header = [], [], []
    for wb_name, g in meta.sort_values(["waterbody", "site"]).groupby("waterbody", sort=True):
        rows.append(None)
        labels.append(wb_name)
        is_header.append(True)
        for sid, r in g.iterrows():
            rows.append(sid)
            labels.append(short_name(r.site, wb_name))
            is_header.append(False)

    code = {"good": 1, "advisory": 2, "poor": 3}
    M = np.full((len(rows), len(years)), np.nan)
    small = []
    for i, sid in enumerate(rows):
        if sid is None:
            continue
        for j, y in enumerate(years):
            gm = wide.loc[sid, y]
            if np.isfinite(gm):
                M[i, j] = code[category(gm)]
                if nwide.loc[sid, y] < MIN_SAMPLES_FOR_GEOMEAN:
                    small.append((j, i))
            else:
                M[i, j] = 0  # not sampled that season

    cmap = ListedColormap(["#f0efec", COLOR_GOOD, COLOR_ADVISORY, COLOR_POOR])
    fig_h = len(rows) * 0.17 + 1.6
    fig, ax = plt.subplots(figsize=(10.5, fig_h))
    ax.pcolormesh(np.arange(len(years) + 1), np.arange(len(rows) + 1), np.ma.masked_invalid(M),
                  cmap=cmap, vmin=-0.5, vmax=3.5, edgecolors=SURFACE, linewidth=1.6)
    if small:
        sx, sy_ = zip(*small)
        ax.scatter(np.array(sx) + 0.5, np.array(sy_) + 0.5, s=6, color=INK, alpha=0.6, lw=0, zorder=3)
    ax.set_ylim(len(rows), 0)
    ax.set_xlim(0, len(years))
    ax.set_yticks(np.arange(len(rows)) + 0.5)
    ax.set_yticklabels(labels, fontsize=7.2)
    for t, h in zip(ax.get_yticklabels(), is_header):
        if h:
            t.set_fontweight("bold")
            t.set_color(INK)
            t.set_fontsize(8)
    ax.tick_params(axis="y", length=0, pad=4)
    ax.set_xticks(np.arange(len(years)) + 0.5)
    ax.set_xticklabels(years, fontsize=8)
    ax.xaxis.tick_top()
    ax.tick_params(axis="x", length=0, pad=3)
    ax.grid(False)
    for sp in ax.spines.values():
        sp.set_visible(False)

    handles = [Patch(color=COLOR_GOOD, label=CATEGORY_LABELS["good"]),
               Patch(color=COLOR_ADVISORY, label=CATEGORY_LABELS["advisory"]),
               Patch(color=COLOR_POOR, label=CATEGORY_LABELS["poor"]),
               Patch(color="#f0efec", label="Not sampled"),
               Line2D([], [], marker="o", ls="", color=INK, alpha=0.6, ms=3,
                      label=f"Fewer than {MIN_SAMPLES_FOR_GEOMEAN} samples")]
    ax.legend(handles=handles, loc="lower left", bbox_to_anchor=(0, 1),
              bbox_transform=offset_copy(ax.transAxes, fig=fig, y=20, units="points"),
              ncols=5, fontsize=8, handlelength=1.2, columnspacing=1.2, borderaxespad=0)
    ax.annotate("Sites grouped by waterbody. Color = that season's geomean (MPN / 100 ml), "
                "regular-season samples only.", xy=(0, 1), xycoords="axes fraction",
                xytext=(0, 44), textcoords="offset points", fontsize=9, color=INK_2)
    ax.set_title("Seasonal geometric mean at every site", pad=62)
    fig.savefig(out)
    plt.close(fig)


# ---- 3. Rainfall effect -----------------------------------------------------

RAIN_EDGES = [0, 0.005, 0.25, 0.5, 1.0, np.inf]   # inches; bins include the left edge
RAIN_LABELS = ["None\n(0 or trace)", "0.01–0.24\"", "0.25–0.49\"", "0.50–0.99\"", "1\" or more"]


def chart_rainfall(season, out: Path, tables: Path):
    s = season.dropna(subset=["rain_in"]).copy()
    s["bin"] = pd.cut(s.rain_in, RAIN_EDGES, right=False, labels=RAIN_LABELS)
    rows = []
    for lab in RAIN_LABELS:
        g = s[s.bin == lab].mpn
        rows.append({"rain_bin": lab.replace("\n", " "), "n_samples": len(g), "geomean": geomean(g),
                     "q25": g.quantile(0.25), "median": g.median(), "q75": g.quantile(0.75),
                     "pct_over_limit": 100 * (g > SINGLE_SAMPLE_LIMIT).mean()})
    t = pd.DataFrame(rows)
    _r(t).to_csv(tables / "rainfall_effect.csv", index=False)

    fig, (a1, a2) = plt.subplots(1, 2, figsize=(11.5, 4.9), gridspec_kw={"wspace": 0.42})
    x = np.arange(len(RAIN_LABELS))
    _mpn_axis(a1)
    a1.set_ylim(MPN_FLOOR, 3000)
    a1.set_xlim(-0.5, len(RAIN_LABELS) - 0.5)
    _quality_bands(a1)
    a1.vlines(x, t.q25, t.q75, color=SERIES, lw=2.5, alpha=0.35, zorder=2, capstyle="round")
    a1.plot(x, t.geomean, ls="", marker="o", ms=8, color=SERIES, mec=SURFACE, mew=2, zorder=3)
    for xi, gm in zip(x, t.geomean):
        a1.annotate(f"{gm:,.0f}", (xi, gm), xytext=(9, 0), textcoords="offset points",
                    va="center", fontsize=8.5, color=INK,
                    bbox=dict(boxstyle="square,pad=0.15", fc=SURFACE, ec="none", alpha=0.8))
    a1.set_xticks(x)
    a1.set_xticklabels(RAIN_LABELS, fontsize=8.5)
    a1.grid(axis="x", visible=False)
    a1.set_ylabel("MPN / 100 ml (log scale)")
    a1.set_xlabel(f"Rain at Central Park, sampling day + {RAIN_DAYS - 1} days before")
    a1.set_title("Bacteria levels rise with rainfall", pad=28)
    _subtitle(a1, "Dot = geometric mean; line = middle 50% of samples")

    bars = a2.bar(x, t.pct_over_limit, width=0.55, color=SERIES, zorder=2)
    for b, v in zip(bars, t.pct_over_limit):
        a2.annotate(f"{v:.0f}%", (b.get_x() + b.get_width() / 2, v), xytext=(0, 3),
                    textcoords="offset points", ha="center", va="bottom", fontsize=9, color=INK)
    a2.set_xticks(x)
    a2.set_xticklabels([f"{l}\nn={n:,}" for l, n in zip(RAIN_LABELS, t.n_samples)], fontsize=8.5)
    a2.set_ylim(0, 100)
    a2.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:.0f}%"))
    a2.grid(axis="x", visible=False)
    a2.set_xlabel(f"Rain at Central Park, sampling day + {RAIN_DAYS - 1} days before")
    a2.set_title(f"Share of samples over {SINGLE_SAMPLE_LIMIT} MPN", pad=28)
    _subtitle(a2, f"All sites, {int(s.year.min())}–{int(s.year.max())} seasons with rainfall data "
                  f"({len(s):,} samples)")
    fig.savefig(out)
    plt.close(fig)


# ---- 4. Waterbody comparison ------------------------------------------------

def _right_col(ax, y, text, x=1.01, **kw):
    kw.setdefault("fontsize", 8)
    kw.setdefault("color", INK_2)
    ax.annotate(text, (x, y), xycoords=("axes fraction", "data"), va="center", ha="left", **kw)


def chart_waterbodies(season, out: Path, tables: Path):
    last = int(season.year.max())
    first = last - RECENT_SEASONS + 1
    s = season[season.year >= first]
    site_gm = s.groupby(["waterbody", "site_id"]).mpn.agg(gm=geomean, n="size").reset_index()
    site_gm = site_gm[site_gm.n >= MIN_SAMPLES_FOR_GEOMEAN]
    wb = s.groupby("waterbody").agg(geomean=("mpn", geomean), n_samples=("mpn", "size"),
                                    n_sites=("site_id", "nunique"),
                                    pct_over=("mpn", lambda v: 100 * (v > SINGLE_SAMPLE_LIMIT).mean()))
    wb = wb[wb.n_samples >= 10].sort_values("geomean")
    _r(wb).to_csv(tables / f"waterbody_geomean_{first}_{last}.csv")

    fig_h = 0.3 * len(wb) + 1.6
    fig, ax = plt.subplots(figsize=(10, fig_h))
    _mpn_axis(ax, "x")
    ax.set_xlim(MPN_FLOOR, 5000)
    ax.set_ylim(len(wb) - 0.4, -0.6)
    _quality_bands(ax, axis="x")
    ypos = {w: i for i, w in enumerate(wb.index)}
    sg = site_gm[site_gm.waterbody.isin(wb.index)]
    ax.scatter(sg.gm, sg.waterbody.map(ypos), s=22, color=SERIES, alpha=0.35, lw=0, zorder=2)
    ax.scatter(wb.geomean, np.arange(len(wb)), s=70, color=SERIES, edgecolor=SURFACE, lw=2, zorder=3)
    _right_col(ax, -1.0, "Geomean", fontsize=8, color=INK, fontweight="bold")
    _right_col(ax, -1.0, "Sites · samples", x=1.105, fontsize=8, color=INK, fontweight="bold")
    for i, (w, r) in enumerate(wb.iterrows()):
        _right_col(ax, i, f"{r.geomean:,.0f}", color=INK, fontsize=8.5)
        _right_col(ax, i, f"{int(r.n_sites)} · {int(r.n_samples):,}", x=1.105)
    ax.set_yticks(np.arange(len(wb)))
    ax.set_yticklabels(wb.index, fontsize=8.5)
    ax.tick_params(axis="y", length=0)
    ax.grid(axis="y", visible=False)
    ax.set_xlabel("Geometric mean, MPN / 100 ml (log scale)")
    ax.set_title(f"Recent water quality by waterbody, {first}–{last} seasons", pad=44)
    ax.annotate(f"Large dot = all samples in the waterbody; small dots = individual sites "
                f"(≥{MIN_SAMPLES_FOR_GEOMEAN} samples). Cleanest at top; waterbodies with "
                f"<10 samples left out.", xy=(0, 1), xycoords="axes fraction", xytext=(0, 24),
                textcoords="offset points", fontsize=8.8, color=INK_2)
    fig.savefig(out)
    plt.close(fig)


# ---- 5. Latest 30-day status -----------------------------------------------

def chart_latest_window(win, sites, out: Path):
    w = win.merge(sites[["site_id", "site", "waterbody"]], on="site_id").sort_values("geomean")
    start, end = win.window_start.iloc[0], win.window_end.iloc[0]
    fig_h = 0.2 * len(w) + 1.6
    fig, ax = plt.subplots(figsize=(10, fig_h))
    _mpn_axis(ax, "x")
    y = np.arange(len(w))
    ax.set_ylim(len(w) - 0.4, -0.8)
    _quality_bands(ax, axis="x")
    full = (w.n >= MIN_SAMPLES_FOR_GEOMEAN).values
    ax.hlines(y, MPN_FLOOR, w.geomean, color=GRID, lw=0.8, zorder=1)
    ax.scatter(w.geomean[full], y[full], s=40, color=SERIES, edgecolor=SURFACE, lw=1.5, zorder=3)
    ax.scatter(w.geomean[~full], y[~full], s=40, facecolor=SURFACE, edgecolor=SERIES, lw=1.5, zorder=3)
    _right_col(ax, -1.2, "Geomean", color=INK, fontweight="bold")
    _right_col(ax, -1.2, "Samples", x=1.1, color=INK, fontweight="bold")
    for yi, (_, r) in zip(y, w.iterrows()):
        _right_col(ax, yi, f"{r.geomean:,.0f}", color=INK, fontsize=7.5)
        _right_col(ax, yi, f"{r.n}", x=1.1, fontsize=7.5)
    ax.set_yticks(y)
    ax.set_yticklabels(w.site, fontsize=7.2)
    ax.tick_params(axis="y", length=0)
    ax.grid(axis="y", visible=False)
    ax.set_xlabel("30-day geometric mean, MPN / 100 ml (log scale)")
    ax.set_title(f"30-day geometric mean, {start:%b %-d} – {end:%b %-d, %Y}", pad=44)
    ax.annotate(f"Same calculation as the GeoMean tabs. Cleanest at top. Hollow dot = fewer than "
                f"{MIN_SAMPLES_FOR_GEOMEAN} samples in the window.", xy=(0, 1),
                xycoords="axes fraction", xytext=(0, 24), textcoords="offset points",
                fontsize=8.8, color=INK_2)
    fig.savefig(out)
    plt.close(fig)


# ---- 6. Per-site time series PDF -------------------------------------------

def _season_x(dates: pd.Series, seasons: list[int]) -> pd.Series:
    """Place each date inside its season's slot: season i spans x = i .. i+1."""
    idx = {y: i for i, y in enumerate(seasons)}
    yr = dates.dt.year
    st = pd.to_datetime(dict(year=yr, month=SEASON_START[0], day=SEASON_START[1]))
    en = pd.to_datetime(dict(year=yr, month=SEASON_END[0], day=SEASON_END[1]))
    frac = ((dates - st) / (en - st)).clip(0, 1)
    return yr.map(idx) + 0.07 + 0.86 * frac


def site_timeseries_pdf(samples, sites, out: Path, png_dir: Path | None = None):
    order = sites.sort_values(["waterbody", "site"])
    with PdfPages(out) as pdf:
        for _, srow in order.iterrows():
            allg = samples[samples.site_id == srow.site_id]
            g = allg[allg.in_season]
            if g.empty:
                continue
            seasons = sorted(int(y) for y in g.year.unique())
            gx = g.assign(x=_season_x(g.date, seasons))

            fig, ax = plt.subplots(figsize=(11, 4.8))
            _mpn_axis(ax)
            ax.set_xlim(0, len(seasons))
            _quality_bands(ax)
            for i in range(1, len(seasons)):
                ax.axvline(i, color=GRID, lw=0.8, zorder=1)
            dry, wet, unk = gx[gx.wet == 0], gx[gx.wet == 1], gx[gx.wet.isna()]
            ax.scatter(dry.x, dry.mpn, s=20, color=SERIES, lw=0, alpha=0.85, zorder=3,
                       label=f"Dry weather (<{WET_WEATHER_INCHES}\" rain)")
            ax.scatter(wet.x, wet.mpn, s=20, color=SERIES_2, lw=0, alpha=0.85, zorder=3,
                       label=f"Wet weather (≥{WET_WEATHER_INCHES}\" rain)")
            if len(unk):
                ax.scatter(unk.x, unk.mpn, s=20, facecolor="none", edgecolor=INK_2, lw=0.9,
                           zorder=3, label="Rainfall not entered yet")
            ticklabels = []
            for i, y in enumerate(seasons):
                gy = gx[gx.year == y]
                gm = geomean(gy.mpn)
                ax.hlines(gm, i + 0.05, i + 0.95, color=INK, lw=2, zorder=4, capstyle="round")
                ticklabels.append(f"{y}\n{gm:,.0f}")
            ax.set_xticks(np.arange(len(seasons)) + 0.5)
            ax.set_xticklabels(ticklabels, fontsize=8.5 if len(seasons) < 12 else 7.8)
            ax.tick_params(axis="x", length=0)
            ax.grid(axis="x", visible=False)
            ax.set_xlabel("Season  (number underneath = that season's geometric mean)", fontsize=8.5)
            ax.set_ylabel("Enterococcus, MPN / 100 ml (log)")
            ax.set_title(srow.site, pad=28)
            st = summarize(g)
            cls = f" · DEC class {srow.dec_class}" if srow.dec_class else ""
            off = len(allg) - len(g)
            off_txt = f" · {off} off-season sample{'s' if off != 1 else ''} not shown" if off else ""
            _subtitle(ax, f"{srow.waterbody} · {srow.borough}{cls} · {st.n:,} samples · overall "
                          f"geomean {st.geomean:,.0f} · {st.pct_over_limit:.0f}% over "
                          f"{SINGLE_SAMPLE_LIMIT}{off_txt}")
            handles, labels_ = ax.get_legend_handles_labels()
            handles.append(Line2D([], [], color=INK, lw=2))
            labels_.append("Season geometric mean")
            ax.legend(handles, labels_, loc="upper left", bbox_to_anchor=(0, -0.2), ncols=4,
                      fontsize=8)
            pdf.savefig(fig)
            if png_dir is not None:
                fn = re.sub(r"[^A-Za-z0-9]+", "_", f"{srow.site_id:03d}_{srow.site}").strip("_")
                fig.savefig(png_dir / f"{fn}.png")
            plt.close(fig)


# =============================================================================
# INTERACTIVE MAP (Leaflet)
# =============================================================================
# The map is one self-contained HTML file: the numbers are embedded as a compact
# JSON table and a small script builds the markers and popups in the browser.
# Only the Leaflet library and the basemap tiles load from the internet.

# Light gray basemap (free, no API key). Swap for any XYZ tile URL you prefer,
# e.g. OpenStreetMap: "https://tile.openstreetmap.org/{z}/{x}/{y}.png"
MAP_TILES = ("https://server.arcgisonline.com/ArcGIS/rest/services/Canvas/"
             "World_Light_Gray_Base/MapServer/tile/{z}/{y}/{x}")
MAP_TILES_ATTR = "Tiles &copy; Esri &mdash; Esri, HERE, Garmin, &copy; OpenStreetMap contributors"
LEAFLET_VERSION = "1.9.4"


def _num(v, nd=1):
    """JSON-safe number: NaN/None -> None, otherwise rounded."""
    if v is None:
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(f):
        return None
    f = round(f, nd)
    return int(f) if nd == 0 or f.is_integer() else f


def _stat_rows(stats: pd.DataFrame, site_ids: set) -> list:
    """One compact row per site:
    [site_id, geomean, n, pct_over, gm_dry, n_dry, gm_wet, n_wet, first, last, last_mpn]"""
    rows = []
    for _, st in stats.sort_values("geomean").iterrows():
        if st.site_id not in site_ids:
            continue
        rows.append([int(st.site_id), _num(st.geomean), int(st.n), _num(st.pct_over_limit, 0),
                     _num(st.geomean_dry), int(st.n_dry), _num(st.geomean_wet), int(st.n_wet),
                     f"{st['first']:%Y-%m-%d}", f"{st['last']:%Y-%m-%d}", _num(st.last_mpn)])
    return rows


def build_map(samples, sites, site_all, site_year, win, out: Path):
    have = sites[sites.site_id.isin(samples.site_id)].dropna(subset=["lat", "lon"])
    ids = set(have.site_id)
    y0, y1 = int(samples.year.min()), int(samples.year.max())

    # Layers in radio-button order; the first one is shown when the map opens.
    seasons = sorted((int(y) for y in site_year.year.unique()), reverse=True)
    layers = [{"name": f"{seasons[0]} season", "label": f"{seasons[0]} season",
               "rows": _stat_rows(site_year[site_year.year == seasons[0]], ids)}]
    if win is not None and len(win):
        s0, s1 = win.window_start.iloc[0], win.window_end.iloc[0]
        layers.append({"name": f"Latest 30 days ({s0:%b %-d}–{s1:%b %-d})",
                       "label": f"30 days, {s0:%b %-d}–{s1:%b %-d, %Y}",
                       "rows": _stat_rows(win, ids)})
    layers.append({"name": f"All seasons ({y0}–{y1})", "label": f"all seasons {y0}–{y1}",
                   "rows": _stat_rows(site_all, ids)})
    for y in seasons[1:]:
        layers.append({"name": f"{y} season", "label": f"{y} season",
                       "rows": _stat_rows(site_year[site_year.year == y], ids)})

    data = {
        # site_id -> [name, waterbody, borough, DEC class, lat, lon]
        "sites": {int(r.site_id): [r.site, r.waterbody, str(r.borough) if pd.notna(r.borough) else "",
                                   r.dec_class, round(float(r.lat), 5), round(float(r.lon), 5)]
                  for r in have.itertuples()},
        # site_id -> [[season, geomean], ...] for the popup trend chart
        "trend": {int(sid): [[int(y), _num(gm)] for y, gm in zip(g.year, g.geomean) if np.isfinite(gm)]
                  for sid, g in site_year.sort_values("year").groupby("site_id")},
        "overall": {int(r.site_id): _num(r.geomean) for r in site_all.itertuples()},
        "layers": layers,
        "years": [y0, y1],
        "cfg": {"good": GOOD_MAX, "adv": ADVISORY_MAX, "limit": SINGLE_SAMPLE_LIMIT,
                "minN": MIN_SAMPLES_FOR_GEOMEAN, "floor": MPN_FLOOR,
                "crit": {k: list(v) for k, v in DEC_CRITERIA.items()},
                "colors": CATEGORY_COLORS, "labels": CATEGORY_LABELS,
                "series": SERIES, "surface": SURFACE, "ink": INK, "ink2": INK_2},
        "bounds": [[float(have.lat.min()), float(have.lon.min())],
                   [float(have.lat.max()), float(have.lon.max())]],
        "title": f"CWQT water quality · {len(have)} sites · {len(samples):,} regular-season "
                 f"samples, {y0}–{y1}",
        "updated": f"Data through {samples.date.max():%b %-d, %Y}",
    }
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")

    page = (_MAP_TEMPLATE
            .replace("__LEAFLET__", LEAFLET_VERSION)
            .replace("__TILES__", MAP_TILES)
            .replace("__ATTR__", MAP_TILES_ATTR)
            .replace("__SURFACE__", SURFACE).replace("__INK2__", INK_2).replace("__INK__", INK)
            .replace("__DATA__", payload))
    out.write_text(page, encoding="utf-8")


_MAP_TEMPLATE = r"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>CWQT water quality map</title>
<link rel="stylesheet" href="https://unpkg.com/leaflet@__LEAFLET__/dist/leaflet.css">
<script src="https://unpkg.com/leaflet@__LEAFLET__/dist/leaflet.js"></script>
<style>
  html, body, #map { height: 100%; margin: 0; }
  body { background: #e5e5e3; font: 12px system-ui, -apple-system, "Segoe UI", sans-serif; color: __INK__; }
  .panel { background: __SURFACE__; border-radius: 8px; box-shadow: 0 1px 4px rgba(0,0,0,.25); }
  #title { position: fixed; top: 12px; left: 56px; z-index: 1000; padding: 8px 12px; font-weight: 600; font-size: 14px; }
  #title small { display: block; font-weight: 400; font-size: 11px; color: __INK2__; margin-top: 2px; }
  #legend { position: fixed; bottom: 24px; left: 12px; z-index: 1000; padding: 10px 12px; max-width: 250px; }
  #legend .row { margin: 2px 0; }
  .dot { display: inline-block; width: 11px; height: 11px; border-radius: 50%; margin-right: 7px; vertical-align: -1px; }
  .muted { color: __INK2__; }
  .layers { padding: 8px 12px; max-height: calc(100vh - 120px); overflow-y: auto; }
  .layers b { display: block; margin-bottom: 3px; }
  .layers label { display: block; line-height: 1.5; cursor: pointer; white-space: nowrap; }
  .layers input { margin: 0 6px 0 0; vertical-align: -1px; }
  .leaflet-popup-content { margin: 12px 14px; font: 12px system-ui, -apple-system, "Segoe UI", sans-serif; }
  .pop { width: 290px; }
  .pop .name { font-weight: 700; font-size: 13px; margin-bottom: 2px; }
  .pop table { border-collapse: collapse; width: 100%; font-size: 12px; }
  .pop td:first-child { padding-right: 10px; }
</style>
</head>
<body>
<div id="map"></div>
<div id="title" class="panel"></div>
<div id="legend" class="panel"></div>
<script id="cwqt-data" type="application/json">__DATA__</script>
<script>
(function () {
  const D = JSON.parse(document.getElementById("cwqt-data").textContent);
  const C = D.cfg;
  const esc = s => String(s).replace(/[&<>"']/g, c => ({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[c]));
  const fmt = v => (v === null || v === undefined) ? "–" : Math.round(v).toLocaleString("en-US");
  const MON = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"];
  const ymd = s => s.split("-").map(Number);
  const dLong = s => { const [y, m, d] = ymd(s); return `${MON[m-1]} ${d}, ${y}`; };
  const span = (a, b) => {
    const [ya, ma, da] = ymd(a), [yb, mb, db] = ymd(b);
    if (a === b) return dLong(a);
    if (ya === yb) return `${MON[ma-1]} ${da} – ${MON[mb-1]} ${db}, ${yb}`;
    return `${MON[ma-1]} ${ya} – ${MON[mb-1]} ${yb}`;
  };
  const category = gm => gm === null ? "nodata" : gm <= C.good ? "good" : gm <= C.adv ? "advisory" : "poor";

  // ---- popup trend chart (inline SVG, log scale) ----
  function sparkline(pts) {
    const w = 250, h = 74, pl = 4, pr = 4, pt = 6, pb = 14, [y0, y1] = D.years;
    const lo = Math.log10(C.floor), hi = Math.log10(5000);
    const X = y => pl + (w - pl - pr) * ((y - y0) / Math.max(1, y1 - y0));
    const Y = v => { v = Math.min(Math.max(v, C.floor), 5000);
                     return pt + (h - pt - pb) * (1 - (Math.log10(v) - lo) / (hi - lo)); };
    let s = `<svg width="${w}" height="${h}" viewBox="0 0 ${w} ${h}" role="img" aria-label="Seasonal geomean trend" style="display:block">`;
    [[C.floor, C.good, "good"], [C.good, C.adv, "advisory"], [C.adv, 5000, "poor"]].forEach(([a, b, k]) => {
      s += `<rect x="${pl}" y="${Y(b).toFixed(1)}" width="${w - pl - pr}" height="${(Y(a) - Y(b)).toFixed(1)}" fill="${C.colors[k]}" fill-opacity="0.10"/>`;
    });
    for (let i = 1; i < pts.length; i++) {
      if (pts[i][0] - pts[i-1][0] === 1) {
        s += `<line x1="${X(pts[i-1][0]).toFixed(1)}" y1="${Y(pts[i-1][1]).toFixed(1)}" x2="${X(pts[i][0]).toFixed(1)}" y2="${Y(pts[i][1]).toFixed(1)}" stroke="${C.series}" stroke-width="2" stroke-linecap="round"/>`;
      }
    }
    pts.forEach(([y, v]) => {
      s += `<circle cx="${X(y).toFixed(1)}" cy="${Y(v).toFixed(1)}" r="3.5" fill="${C.series}" stroke="#fff" stroke-width="1.5"><title>${y}: ${fmt(v)}</title></circle>`;
    });
    s += `<text x="${pl}" y="${h - 2}" font-size="10" fill="${C.ink2}">${y0}</text>`;
    s += `<text x="${w - pr}" y="${h - 2}" font-size="10" fill="${C.ink2}" text-anchor="end">${y1}</text></svg>`;
    return s;
  }

  function popup(r, layerLabel) {
    const [id, gm, n, pct, gmd, nd, gmw, nw, first, last, lastMpn] = r;
    const [name, wb, boro, cls] = D.sites[id];
    const cat = category(gm);
    const crit = C.crit[cls];
    let critRow = "";
    if (crit) critRow = `<tr><td>NYS DEC class ${esc(cls)} criteria</td><td>geomean ≤ ${crit[0]}, ≤10% of samples over ${crit[1]}</td></tr>`;
    else if (cls) critRow = `<tr><td>NYS DEC class</td><td>${esc(cls)}</td></tr>`;
    const wetdry = (nd || nw) ? `<tr><td>Geomean, dry / wet weather</td><td>${fmt(gmd)} (n=${nd}) / ${fmt(gmw)} (n=${nw})</td></tr>` : "";
    const few = n < C.minN ? `<div class="muted" style="margin-top:4px">Only ${n} sample${n !== 1 ? "s" : ""}; a geomean needs ≥${C.minN} to compare with standards.</div>` : "";
    return `<div class="pop">
      <div class="name">${esc(name)}</div>
      <div class="muted" style="margin-bottom:6px">${esc(wb)} · ${esc(boro)} · Site ID ${id}</div>
      <div style="margin-bottom:6px"><span class="dot" style="width:10px;height:10px;margin-right:6px;background:${C.colors[cat]}"></span><b>${esc(C.labels[cat])}</b> <span class="muted">(${esc(layerLabel)})</span></div>
      <table>
        <tr><td>Geometric mean</td><td><b>${fmt(gm)}</b> MPN/100 ml</td></tr>
        <tr><td>Samples</td><td>${n.toLocaleString("en-US")} &nbsp;<span class="muted">${span(first, last)}</span></td></tr>
        <tr><td>Samples over ${C.limit}</td><td>${pct}%</td></tr>
        ${wetdry}
        <tr><td>Most recent result</td><td>${fmt(lastMpn)} on ${dLong(last)}</td></tr>
        ${critRow}
      </table>${few}
      <div class="muted" style="margin-top:8px">Seasonal geomean, all years (overall ${fmt(D.overall[id])}):</div>
      ${sparkline(D.trend[id] || [])}
    </div>`;
  }

  // ---- map ----
  const map = L.map("map", { zoomControl: true, preferCanvas: false });
  L.tileLayer("__TILES__", { attribution: "__ATTR__", maxZoom: 16 }).addTo(map);
  L.control.scale().addTo(map);
  map.fitBounds(D.bounds);

  const groups = D.layers.map(() => null);
  function buildGroup(i) {
    const layer = D.layers[i], g = L.layerGroup();
    layer.rows.forEach(r => {
      const [id, gm, n] = r, s = D.sites[id];
      const cat = category(gm), few = n < C.minN;
      L.circleMarker([s[4], s[5]], {
        radius: few ? 6 : 7, color: few ? C.ink2 : C.surface, weight: few ? 1.5 : 2,
        dashArray: few ? "3,3" : null, fillColor: C.colors[cat], fillOpacity: few ? 0.55 : 0.95
      })
        .bindTooltip(`${esc(s[0])} — geomean ${fmt(gm)} (${esc(C.labels[cat].split(":")[0])}, n=${n})`, { sticky: true })
        .bindPopup(() => popup(r, layer.label), { maxWidth: 320 })
        .addTo(g);
    });
    return g;
  }
  let current = -1;
  function show(i) {
    if (current >= 0) map.removeLayer(groups[current]);
    groups[i] = groups[i] || buildGroup(i);
    groups[i].addTo(map);
    current = i;
  }

  const Ctl = L.Control.extend({
    options: { position: "topright" },
    onAdd() {
      const div = L.DomUtil.create("div", "panel layers");
      div.innerHTML = "<b>Show results for</b>" + D.layers.map((l, i) =>
        `<label><input type="radio" name="cwqt-layer" value="${i}"${i === 0 ? " checked" : ""}>${esc(l.name)}</label>`).join("");
      L.DomEvent.disableClickPropagation(div);
      L.DomEvent.disableScrollPropagation(div);
      div.addEventListener("change", e => show(Number(e.target.value)));
      return div;
    }
  });
  map.addControl(new Ctl());
  show(0);

  document.getElementById("title").innerHTML = `${esc(D.title)}<small>${esc(D.updated)}</small>`;
  document.getElementById("legend").innerHTML =
    `<div style="font-weight:700;margin-bottom:4px">Enterococcus geometric mean</div>
     <div class="muted" style="margin-bottom:6px">MPN per 100 ml, NYC DOH primary-contact scale</div>` +
    ["good", "advisory", "poor"].map(k => `<div class="row"><span class="dot" style="background:${C.colors[k]}"></span>${esc(C.labels[k])}</div>`).join("") +
    `<div class="row" style="margin-top:4px"><span class="dot" style="width:9px;height:9px;border:1.5px dashed ${C.ink2};opacity:.8"></span>Fewer than ${C.minN} samples</div>
     <div class="muted" style="margin-top:6px">Pick a season at top right. Click a site for details and its trend.</div>`;
})();
</script>
</body>
</html>
"""


# =============================================================================
# MAIN
# =============================================================================

def download_sheet(sheet_id: str, dest: Path) -> Path:
    """Download the Google Sheet as .xlsx (formulas' last calculated values included).

    Two ways to give access:
      * Share the sheet as "Anyone with the link can view" -> no credentials needed.
      * Keep it private and share it with a Google service account; put that
        account's JSON key in the GOOGLE_SERVICE_ACCOUNT_JSON environment variable
        (needs `pip install google-auth requests`).
    """
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    headers = {"User-Agent": "cwqt-visualize"}
    key = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON", "").strip()
    if key:
        from google.auth.transport.requests import Request as GRequest
        from google.oauth2 import service_account
        creds = service_account.Credentials.from_service_account_info(
            json.loads(key), scopes=["https://www.googleapis.com/auth/drive.readonly"])
        creds.refresh(GRequest())
        headers["Authorization"] = f"Bearer {creds.token}"
    try:
        with urlopen(Request(url, headers=headers), timeout=120) as r:
            body = r.read()
    except (HTTPError, URLError) as e:
        sys.exit(f"Could not download the Google Sheet ({e}). Check the sheet ID and sharing.")
    if not body.startswith(b"PK"):  # .xlsx files are zip archives; anything else is a login page
        sys.exit("Google returned a web page instead of a spreadsheet. The sheet is probably not "
                 "shared with link viewers or with the service account.")
    dest.write_bytes(body)
    print(f"Downloaded sheet ({len(body) / 1e6:.1f} MB) to {dest}")
    return dest


def write_qa(qa: list[QAIssue], checks: list[str], out: Path):
    lines = ["CWQT data quality report", "=" * 24, ""]
    lines += ["Cross-check against the workbook's own GeoMean tabs:"] + [f"  {c}" for c in checks] + [""]
    for cat in dict.fromkeys(q.category for q in qa):
        lines.append(cat)
        lines.append("-" * len(cat))
        lines += [f"  - {q.detail}" for q in qa if q.category == cat]
        lines.append("")
    out.write_text("\n".join(lines), encoding="utf-8")


def _r(df: pd.DataFrame) -> pd.DataFrame:
    """Round numeric columns for CSV output; format dates as YYYY-MM-DD."""
    df = df.copy()
    for c in df.columns:
        if pd.api.types.is_datetime64_any_dtype(df[c]):
            df[c] = df[c].dt.strftime("%Y-%m-%d")
        elif pd.api.types.is_float_dtype(df[c]):
            df[c] = df[c].round(2)
    return df


def main(argv=None):
    ap = argparse.ArgumentParser(description="Maps and charts from the CWQT master data sheet.")
    ap.add_argument("xlsx", type=Path, nargs="?",
                    help="CWQT master data workbook (.xlsx). Leave out to download it with --sheet-id.")
    ap.add_argument("--sheet-id", default=os.environ.get("CWQT_SHEET_ID"),
                    help="Google Sheet ID to download (the long code in the sheet's URL). "
                         "Can also be set with the CWQT_SHEET_ID environment variable.")
    ap.add_argument("--out", type=Path, default=Path("cwqt_outputs"), help="output folder")
    ap.add_argument("--window-end", type=str, default=None,
                    help="last day of the 30-day window (YYYY-MM-DD); default = latest sample date")
    ap.add_argument("--average-splits", action="store_true",
                    help="average lab/field split results recorded in Notes (per the FAQ). "
                         "Off by default so numbers match the GeoMean tabs.")
    ap.add_argument("--site-pngs", action="store_true",
                    help="also save one PNG per site (in addition to the PDF)")
    ap.add_argument("--no-map", action="store_true", help="skip the interactive map")
    args = ap.parse_args(argv)

    out = args.out
    tables = out / "tables"
    tables.mkdir(parents=True, exist_ok=True)
    _style()
    if args.xlsx is None:
        if not args.sheet_id:
            ap.error("give an .xlsx file or a --sheet-id to download")
        args.xlsx = download_sheet(args.sheet_id, out / "CWQT_Master_Data_Sheet.xlsx")

    print(f"Reading {args.xlsx} …")
    qa: list[QAIssue] = []
    sites = load_sites(args.xlsx, qa)
    samples = load_samples(args.xlsx, sites, qa, average_splits=args.average_splits)
    print(f"  {len(samples):,} samples at {samples.site_id.nunique()} sites, "
          f"{samples.date.min():%b %Y} – {samples.date.max():%b %Y}")
    checks = [] if args.average_splits else verify_against_workbook(args.xlsx, samples)
    for c in checks:
        print("  " + c)

    season = samples[samples.in_season]          # regular-season samples only
    print(f"  {len(season):,} regular-season samples ({SEASON_START[0]}/{SEASON_START[1]}–"
          f"{SEASON_END[0]}/{SEASON_END[1]}) used for seasonal statistics")
    site_all = site_stats(season)
    site_year = site_stats(season, by_year=True)
    end = pd.Timestamp(args.window_end) if args.window_end else samples.date.max()
    win = rolling_window_stats(samples, end, 30)

    meta = sites[["site_id", "site", "waterbody", "borough", "dec_class", "lat", "lon"]]
    _r(meta.merge(site_all, on="site_id")).to_csv(tables / "site_summary_all_seasons.csv", index=False)
    _r(meta.merge(site_year, on="site_id")).to_csv(tables / "site_summary_by_season.csv", index=False)
    _r(meta.merge(win, on="site_id")).to_csv(tables / "latest_30_day_window.csv", index=False)

    steps = [
        ("1_program_trend.png", lambda p: chart_program_trend(season, site_year, p, tables)),
        ("2_site_by_season_heatmap.png", lambda p: chart_site_heatmap(site_year, sites, p, tables)),
        ("3_rainfall_effect.png", lambda p: chart_rainfall(season, p, tables)),
        ("4_waterbody_comparison.png", lambda p: chart_waterbodies(season, p, tables)),
        ("5_latest_30_day_status.png", lambda p: chart_latest_window(win, sites, p)),
    ]
    for name, fn in steps:
        fn(out / name)
        print(f"  wrote {out / name}")

    png_dir = None
    if args.site_pngs:
        png_dir = out / "site_pngs"
        png_dir.mkdir(exist_ok=True)
    site_timeseries_pdf(samples, sites, out / "site_timeseries.pdf", png_dir)
    print(f"  wrote {out / 'site_timeseries.pdf'}")

    if not args.no_map:
        build_map(season, sites, site_all, site_year, win, out / "cwqt_map.html")
        print(f"  wrote {out / 'cwqt_map.html'} ({(out / 'cwqt_map.html').stat().st_size / 1e3:,.0f} KB)")

    # Small machine-readable summary (used by the automatic update to detect new data)
    (out / "status.json").write_text(json.dumps({
        "data_through": f"{samples.date.max():%Y-%m-%d}",
        "samples": int(len(samples)),
        "sites": int(samples.site_id.nunique()),
    }, indent=2) + "\n", encoding="utf-8")

    write_qa(qa, checks, out / "data_quality_report.txt")
    print(f"  wrote {out / 'data_quality_report.txt'} ({len(qa)} items)")
    print("Done.")


if __name__ == "__main__":
    main()
